"""ATS score for a tailored Word resume against its job description (2026-10-06).

Real ATS products (Workday, Greenhouse, iCIMS...) don't publish a score, so this is a
Jobscan-style estimate on the actual .docx text, out of 100:

  Hard skills   50  JD technology keywords found in the resume (inventory skills and known
                    technologies the JD names, plus the scorer's matched/gap lists)
  Job title     10  the JD's title words appear in the resume (headline)
  JD language   15  the JD's most frequent meaningful terms appear in the resume
  Format        25  5 each: text extracts cleanly; standard sections (Summary, Skills,
                    Experience, Education); contact info (email, phone, LinkedIn); every role
                    has month-year dates; no tables/images/text boxes/header-footer text and
                    <= 2 pages

Target >= 90. Missing hard skills the candidate genuinely lacks cap the score honestly; the
report lists them so they can go in the Sheet notes instead of onto the resume.

    python3 scripts/ats_check.py <resume.docx> <jd.txt> --title "Senior Data Engineer" \
        [--company Acme] [--matched "a, b"] [--gaps "c, d"] [--pages 2]
"""
import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
from inventory import appears_in, build_index, load_inventory, lookup, normalize
from tech_lexicon import TECH_TERMS

TARGET = 90
STOP = set("""a about above across after again against all also an and any are as at be been
before being below between both but by can could did do does doing down during each few for
from further had has have having he her here hers him his how i if in into is it its itself
just me more most my no nor not now of off on once only or other our ours out over own same
she should so some such than that the their theirs them then there these they this those
through to too under until up very was we were what when where which while who whom why will
with would you your yours ability able etc including include includes within across using use
used work working works role team teams experience experienced years year strong excellent
good great new well plus preferred required requirements responsibilities qualifications
must may might per via like help helps support supports ensure ensuring based job position
candidate candidates company opportunity benefits salary pay range equal employer status
applicants apply application individuals regard without race color religion sex sexual
orientation gender identity national origin disability veteran protected law location remote
hybrid onsite office time full part day days week weeks month months type level senior junior
skills skill knowledge understanding familiarity proficiency hands deep solid proven
appropriate applying approved embrace champion curiosity passion passionate improve improving
enhance accelerate drive drives driving deliver delivering make makes making build building
develop developing design designing related relevant modern innovative best practices practice
looking join joining bring play ways where whether will also highly high-performing world
class leading leader leaders across through other others within while each every""".split())
WORD = re.compile(r"[A-Za-z][A-Za-z+#.\-]*[A-Za-z+#]|[A-Za-z]")
MONTH_YEAR = re.compile(r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+(19|20)\d{2}",
                        re.IGNORECASE)
SECTIONS = ["summary", "skills", "experience", "education"]


def docx_text(path):
    """-> (full text, structure problems found)"""
    doc = Document(path)
    problems = []
    if doc.tables:
        problems.append(f"{len(doc.tables)} table(s)")
    body = doc.element.body
    if body.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"):
        problems.append("images or drawings")
    if body.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent"):
        problems.append("text boxes")
    for section in doc.sections:
        for part in (section.header, section.footer):
            if any(p.text.strip() for p in part.paragraphs):
                problems.append("text in header/footer (many ATS skip it)")
                break
    text = "\n".join(p.text for p in doc.paragraphs)
    return text, problems


def jd_tech_keywords(jd, company, matched, gaps):
    """Technology keywords the JD asks for: the scorer's lists plus a fresh scan."""
    text = re.sub(re.escape(company), " ", jd, flags=re.IGNORECASE) if company else jd
    found, seen = [], set()

    def add(term):
        key = normalize(term)
        if key and key not in seen and not any(key in s or s in key for s in seen
                                                if len(s) > 3 and len(key) > 3):
            seen.add(key)
            found.append(term)

    for term in list(matched) + list(gaps):
        add(term)
    items = load_inventory()
    for item in items:
        for term in sorted([item["name"], *item["aliases"]], key=len, reverse=True):
            if appears_in(text, term):
                add(term)
                break
    for term in sorted(TECH_TERMS, key=len, reverse=True):
        if appears_in(text, term):
            add(term)
    return found


def jd_terms(jd, company, title, limit=15):
    words = [w.lower() for w in WORD.findall(jd)]
    skip = STOP | {w.lower() for w in WORD.findall(company or "")}
    counts = Counter(w for w in words if len(w) > 3 and w not in skip)
    return [w for w, _ in counts.most_common(limit)]


def score(docx_path, jd, title, company="", matched=(), gaps=(), pages=None):
    text, structure_problems = docx_text(docx_path)
    lower = text.lower()
    index = build_index(load_inventory())

    keywords = jd_tech_keywords(jd, company, matched, gaps)
    present = [k for k in keywords if appears_in(text, k)]
    missing = [k for k in keywords if k not in present]
    lacking = [k for k in missing if not lookup(k, index)]       # candidate doesn't have these
    fixable = [k for k in missing if lookup(k, index)]           # he has them: add them
    hard = 50 * len(present) / len(keywords) if keywords else 50

    title_words = [w for w in WORD.findall(title.lower())
                   if w not in STOP and not w.isdigit() and w not in {"i", "ii", "iii", "iv", "v",
                                                                       "sr", "jr", "lead",
                                                                       "staff", "principal"}]
    title_hits = [w for w in title_words if w in lower]
    title_pts = 10 * len(title_hits) / len(title_words) if title_words else 10

    terms = jd_terms(jd, company, title)
    term_hits = [t for t in terms if t in lower]
    lang_pts = 15 * len(term_hits) / len(terms) if terms else 15

    fmt, fmt_notes = 0, []
    if len(text.strip()) > 500:
        fmt += 5
    else:
        fmt_notes.append("little text extracted")
    missing_sections = [s for s in SECTIONS if not re.search(rf"^\s*{s}\s*$", lower, re.MULTILINE)]
    fmt += 5 if not missing_sections else 0
    if missing_sections:
        fmt_notes.append(f"missing section headings: {', '.join(missing_sections)}")
    contact = {"email": re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text),
               "phone": re.search(r"\+?\d[\d\s().-]{8,}\d", text),
               "linkedin": re.search(r"linkedin\.com/in/", lower)}
    fmt += 5 if all(contact.values()) else 0
    if not all(contact.values()):
        fmt_notes.append("missing contact: " + ", ".join(k for k, v in contact.items() if not v))
    role_lines = [l for l in text.splitlines() if l.count("|") >= 2 and "\t" in l]
    undated = [l.split("\t")[0] for l in role_lines if not MONTH_YEAR.search(l)]
    fmt += 5 if role_lines and not undated else 0
    if not role_lines or undated:
        fmt_notes.append("roles without month-year dates")
    layout_problems = list(structure_problems)
    if pages is not None and pages > 2:
        layout_problems.append(f"{pages} pages")
    fmt += 5 if not layout_problems else 0
    fmt_notes += layout_problems

    total = round(hard + title_pts + lang_pts + fmt)
    return {
        "score": total,
        "passed": total >= TARGET,
        "hard_skills": f"{len(present)}/{len(keywords)}",
        "missing_you_have": fixable,
        "missing_you_lack": lacking,
        "title_words_missing": [w for w in title_words if w not in title_hits],
        "jd_terms_missing": [t for t in terms if t not in term_hits],
        "format_issues": fmt_notes,
        "parts": {"hard_skills": round(hard, 1), "title": round(title_pts, 1),
                  "jd_language": round(lang_pts, 1), "format": fmt},
    }


def summary_line(result):
    parts = [f"ATS {result['score']}/100", f"hard skills {result['hard_skills']}"]
    if result["missing_you_lack"]:
        parts.append("not claimed (you lack): " + ", ".join(result["missing_you_lack"]))
    if result["format_issues"]:
        parts.append("format: " + "; ".join(result["format_issues"]))
    return " | ".join(parts)


def split_list(text):
    return [k.strip() for k in (text or "").split(",") if k.strip()]


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("docx")
    parser.add_argument("jd", help="text file with the job description")
    parser.add_argument("--title", required=True)
    parser.add_argument("--company", default="")
    parser.add_argument("--matched", default="")
    parser.add_argument("--gaps", default="")
    parser.add_argument("--pages", type=int)
    args = parser.parse_args()
    result = score(args.docx, Path(args.jd).read_text(encoding="utf-8"), args.title,
                   args.company, split_list(args.matched), split_list(args.gaps), args.pages)
    print(json.dumps(result, indent=1))
    print(("PASS " if result["passed"] else "BELOW TARGET ") + summary_line(result))
