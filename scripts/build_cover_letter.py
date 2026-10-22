"""Build a cover letter as a one-page Word file and PDF that match the tailored resume
(2026-10-22, requested by Chidvilas).

    python3 scripts/build_cover_letter.py <letter.md> <out.docx> --resume <resume.md> \
        --company "Deloitte" --title "Forward Deployed Engineer" [--accent "#RRGGBB"] [--pdf]

Input: the letter Claude wrote, from the greeting to the signature, as plain paragraphs
separated by blank lines:

    Dear Hiring Team,

    <opening paragraph>
    ...
    Sincerely,
    Chidvilas Parusha

Output layout, matching the resume (same scripts/build_docx.py helpers, so the same Calibri,
brand color and contact line): name in the company color, the resume's contact line, a thin
colored rule, then the date, the company, a bold "Re: <title>" line, and the letter. Exactly one
page: tries roomier to tighter layouts and keeps the first that LibreOffice puts on one page.
"""
import argparse
import re
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_docx import (FONT, _add_runs, _bottom_border, _set_font, _spacing, measure, parse,
                        readable_accent, to_pdf)
from text_rules import clean_dashes

# (body pt, margin inches), roomiest first. A cover letter reads best with generous margins.
LAYOUTS = [(11, 1.0), (11, 0.85), (10.5, 0.85), (10.5, 0.75), (10, 0.75)]
LETTER_LINE_HEIGHT = 1.2


def letter_paragraphs(letter_md):
    """Blank-line separated paragraphs; single line breaks inside one (the sign-off) are kept."""
    text = clean_dashes(letter_md).strip()
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)          # no bold inside the letter body
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


def build(letter_md, resume_md, out_path, company, title, accent=None, body_pt=11,
          margin_in=1.0, date=None):
    color, hex_color = readable_accent(accent)
    name, header, _ = parse(clean_dashes(resume_md))
    contact = header[1] if len(header) > 1 else ""
    date = date or datetime.now(ZoneInfo("America/New_York"))

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(margin_in)
        section.left_margin = section.right_margin = Inches(margin_in)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(body_pt)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    props = doc.core_properties
    props.title = f"{name} Cover Letter"
    props.author, props.subject = name, f"Cover letter: {title}, {company}"

    p = doc.add_paragraph()
    _spacing(p, size=body_pt + 7)
    _add_runs(p, name, body_pt + 7, color=color, bold=True)
    p = doc.add_paragraph()
    _spacing(p, after=2, size=body_pt)
    _add_runs(p, contact, body_pt - 0.5)
    _bottom_border(p, hex_color)

    def line(text, before=0, after=0, bold=False):
        par = doc.add_paragraph()
        _spacing(par, before=before, after=after, size=body_pt)
        par.paragraph_format.line_spacing = Pt(round(body_pt * LETTER_LINE_HEIGHT, 1))
        _set_font(par.add_run(text), body_pt, bold=bold)
        return par

    line(f"{date:%B} {date.day}, {date:%Y}", before=14)
    line(f"Hiring Team, {company}", before=8)
    line(f"Re: {title}", before=8, after=6, bold=True)

    for para in letter_paragraphs(letter_md):
        par = doc.add_paragraph()
        _spacing(par, after=body_pt * 0.75, size=body_pt)
        par.paragraph_format.line_spacing = Pt(round(body_pt * LETTER_LINE_HEIGHT, 1))
        for i, text in enumerate(para.split("\n")):
            run = par.add_run(("" if i == 0 else "\n") + text.strip())
            _set_font(run, body_pt)
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = False
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return hex_color


def build_fitted(letter_md, resume_md, out_path, company, title, accent=None):
    """Roomiest layout that fits one page. -> (accent hex, pages, body pt)."""
    pages = hex_color = body_pt = None
    for body_pt, margin in LAYOUTS:
        hex_color = build(letter_md, resume_md, out_path, company, title, accent, body_pt, margin)
        pages, _ = measure(out_path, margin)
        if pages is None or pages <= 1:
            break
    return hex_color, pages, body_pt


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("letter")
    parser.add_argument("docx")
    parser.add_argument("--resume", required=True, help="the tailored resume (for name/contact)")
    parser.add_argument("--company", required=True)
    parser.add_argument("--title", required=True)
    parser.add_argument("--accent")
    parser.add_argument("--pdf", action="store_true", help="also write the PDF next to the docx")
    args = parser.parse_args()
    letter = Path(args.letter).read_text(encoding="utf-8")
    resume = Path(args.resume).read_text(encoding="utf-8")
    hex_color, pages, size = build_fitted(letter, resume, args.docx, args.company, args.title,
                                          args.accent)
    print(f"Built {args.docx} ({pages if pages is not None else '?'} page(s), {size}pt {FONT}, "
          f"accent {hex_color})")
    if pages is not None and pages > 1:
        print("TOO LONG: the letter does not fit one page even at 10pt; shorten it")
    if args.pdf:
        print(f"PDF: {to_pdf(args.docx, Path(args.docx).parent)}")
